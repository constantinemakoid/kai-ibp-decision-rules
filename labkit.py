import os
import re
import shutil
import subprocess
import tempfile
import uuid

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
MONO = "Courier New"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

STUDENT = "Макоїд Костянтин Вікторович"
GROUP = "Б-121-23-1-ПІ"

DASH_RE = re.compile(r"[—–‒―]|\s-\s|--")
PRONOUN_RE = re.compile(r"(?<![А-Яа-яІіЇїЄєҐґ'’])(я|мене|мені|мною|ми|нас|нам|нами|мій|моя|моє|мої|мого|моєї|моїх|наш|наша|наше|наші|нашого|нашої|наших)(?![А-Яа-яІіЇїЄєҐґ'’])", re.IGNORECASE)


class Report:
    def __init__(self, kind, number, discipline, topic, teacher, variant=None,
                 student_line=None):
        self.doc = Document()
        self._defaults()
        self.fig_n = 0
        self.tbl_n = 0
        self._title(kind, number, discipline, topic, teacher, variant, student_line)

    def _defaults(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(14)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.first_line_indent = Cm(1.25)
        pf.space_after = Pt(0)
        s = self.doc.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(3)
        s.right_margin = Cm(1.5)

    def _run(self, par, text, bold=False, italic=False, size=14, font=FONT):
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = font
        r.font.size = Pt(size)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        return r

    def center(self, text, bold=False, size=14):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        if text:
            self._run(par, text, bold=bold, size=size)
        return par

    def _title(self, kind, number, discipline, topic, teacher, variant, student_line):
        c = self.center
        c("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ")
        c("Державний університет «Київський авіаційний інститут»")
        c("Факультет комп'ютерних наук та технологій")
        c("Кафедра інженерії програмного забезпечення")
        c("")
        c("")
        c("ЗВІТ", bold=True)
        c(f"з {kind} № {number}")
        c(f"з дисципліни «{discipline}»")
        c(f"на тему: «{topic}»")
        if variant is not None:
            c(f"Варіант № {variant}")
        c("")
        c("")
        c(student_line or f"Виконав студент 4 курсу групи {GROUP}")
        c(STUDENT)
        c("")
        c("Перевірив:")
        c(teacher)
        c("")
        c("")
        c("Київ – 2026")
        self.page_break()

    def page_break(self):
        par = self.doc.add_paragraph()
        par.paragraph_format.first_line_indent = Cm(0)
        par.add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text, new_page=False):
        if new_page:
            self.page_break()
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        self._run(par, text, bold=True, size=16)
        return par

    def h2(self, text):
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(8)
        pf.space_after = Pt(2)
        pf.keep_with_next = True
        self._run(par, text, bold=True, size=14)
        return par

    def p(self, text, indent=True, bold_prefix=None, align="justify"):
        par = self.doc.add_paragraph()
        par.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                         "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        if not indent:
            par.paragraph_format.first_line_indent = Cm(0)
        if bold_prefix:
            self._run(par, bold_prefix, bold=True)
        self._run(par, text)
        return par

    def bullets(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet")
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.left_indent = Cm(1.0)
            if isinstance(it, tuple):
                self._run(par, it[0], bold=True)
                self._run(par, it[1])
            else:
                self._run(par, it)

    def numbered(self, items, start=1):
        for i, it in enumerate(items, start):
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.left_indent = Cm(1.0)
            self._run(par, f"{i}. {it}")

    def qa(self, n, question, answer_paragraphs):
        self.p(f"{n}. {question}", indent=False, bold_prefix=None).runs[0].bold = True
        for a in ([answer_paragraphs] if isinstance(answer_paragraphs, str) else answer_paragraphs):
            self.p(a)

    def fig(self, path, caption, width_cm=16.0):
        self.fig_n += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(path, width=Cm(width_cm))
        cap = self.center(f"Рисунок {self.fig_n}. {caption}", size=12)
        cap.paragraph_format.space_after = Pt(6)
        return self.fig_n

    def table(self, headers, rows, widths_cm=None, caption=None, size=11):
        if caption:
            self.tbl_n += 1
            par = self.doc.add_paragraph()
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.keep_with_next = True
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self._run(par, f"Таблиця {self.tbl_n}. {caption}", size=12)
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.line_spacing = 1.0
            self._run(par, str(h), bold=True, size=size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.line_spacing = 1.0
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                self._run(par, str(v), size=size)
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)
        return t

    def code(self, text, size=9.5):
        for line in text.rstrip("\n").split("\n"):
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.0
            pf.space_after = Pt(0)
            self._run(par, line if line else " ", size=size, font=MONO)
        self.doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)

    def save(self, out_dir, basename):
        os.makedirs(out_dir, exist_ok=True)
        docx_path = os.path.join(out_dir, basename + ".docx")
        self.doc.save(docx_path)
        pdf_path = to_pdf(docx_path, out_dir)
        return docx_path, pdf_path


def to_pdf(src, out_dir):
    prof = os.path.join(tempfile.gettempdir(), "lo_prof_" + uuid.uuid4().hex)
    try:
        subprocess.run([SOFFICE, "--headless", "--norestore",
                        f"-env:UserInstallation=file://{prof}",
                        "--convert-to", "pdf", "--outdir", out_dir, src],
                       check=True, capture_output=True, timeout=300)
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    return os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + ".pdf")


def check_pdf(pdf_path, allow_pronouns=()):
    import pymupdf
    d = pymupdf.open(pdf_path)
    blank, dashes, pron = [], [], []
    for i, page in enumerate(d, 1):
        txt = page.get_text()
        if len(txt.strip()) < 30 and not page.get_images():
            blank.append(i)
        for line in txt.split("\n"):
            if DASH_RE.search(line) and "Київ – 2026" not in line:
                dashes.append((i, line.strip()))
            for m in PRONOUN_RE.finditer(line):
                if not any(a.lower() in line.lower() for a in allow_pronouns):
                    pron.append((i, line.strip()))
    res = {"pages": len(d), "blank_pages": blank, "dash_lines": dashes, "pronoun_lines": pron}
    return res


if __name__ == "__main__":
    import json
    import sys
    print(json.dumps(check_pdf(sys.argv[1]), ensure_ascii=False, indent=1))
